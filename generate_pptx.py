import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen presentation
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    c_navy = RGBColor(15, 23, 42)       # #0f172a
    c_blue = RGBColor(2, 132, 199)      # #0284c7 (Brand Accent)
    c_slate = RGBColor(71, 85, 105)     # #475569
    c_light_bg = RGBColor(248, 250, 252)# #f8fafc
    c_card_bg = RGBColor(255, 255, 255) # White
    c_border = RGBColor(226, 232, 240)  # #e2e8f0
    c_red = RGBColor(220, 38, 38)
    c_green = RGBColor(22, 163, 74)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_light_bg
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="CHEAPPHONE LAB | HACKFRONT INDIA 2026"):
        # Header banner text
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = c_blue

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = c_navy

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative left accent bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.4), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = c_blue
    bar.line.fill.background()

    # Title box
    tbox = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.5), Inches(3.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "CheapPhone Lab"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = c_navy

    p2 = tf1.add_paragraph()
    p2.text = "Test Beyond the Flagship."
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = c_blue
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "A developer testing & diagnostics platform built to uncover, prioritize and reproduce application failures under realistic Indian mobile hardware and telecom conditions."
    p3.font.size = Pt(15)
    p3.font.color.rgb = c_slate
    p3.space_before = Pt(20)

    # Metadata pill
    pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.3), Inches(8.5), Inches(1.0))
    pill.fill.solid()
    pill.fill.fore_color.rgb = c_card_bg
    pill.line.color.rgb = c_border
    ptf = pill.text_frame
    ptf.margin_top = Inches(0.18)
    ptf.margin_left = Inches(0.3)
    pp = ptf.paragraphs[0]
    pp.text = "HackFront India 2026 — Track T1.2: AI & Developer Tools"
    pp.font.size = Pt(13)
    pp.font.bold = True
    pp.font.color.rgb = c_navy
    pp2 = ptf.add_paragraph()
    pp2.text = "Target Track: Neo4j Graph Database Track  |  Prototype Demonstration"
    pp2.font.size = Pt(11)
    pp2.font.color.rgb = c_slate

    # ==========================================
    # SLIDE 2: THE PROBLEM STATEMENT
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Problem: The Flagship Fallacy in Indian App Development")

    col_w = Inches(3.6)
    col_h = Inches(5.1)
    y_pos = Inches(1.6)

    cards = [
        {
            "num": "01",
            "title": "Hardware Reality Gap",
            "subtitle": "Flagship vs. Sub-₹8,000 Phones",
            "points": [
                "Developers test on 8GB-12GB RAM phones over 300Mbps broadband Wi-Fi.",
                "Over 400M Indian users use sub-₹8k phones with 1GB-2GB RAM (Android Go edition).",
                "Aggressive Low Memory Killers (LMK) kill app activities abruptly during transit."
            ]
        },
        {
            "num": "02",
            "title": "Telecom & Storage Bottlenecks",
            "subtitle": "2G/EDGE & Constrained Flash",
            "points": [
                "Commuter travel in tier-2/3 cities causes RTT latency spikes > 1500ms and 30%+ packet loss.",
                "Low remaining internal storage (< 500MB) causes SQLite databases to crash with SQLITE_FULL.",
                "Battery Saver governors throttle CPU clock down to 800MHz little cores, causing ANRs."
            ]
        },
        {
            "num": "03",
            "title": "Silent User Abandonment",
            "subtitle": "Catastrophic Churn",
            "points": [
                "Uncaught OOMs and frozen UIs cause users to uninstall apps immediately without sending bug reports.",
                "Duplicate form & order submissions occur because slow networks lag button debounce logic.",
                "Engineering teams lack tools to model and reproduce these stress realities locally."
            ]
        }
    ]

    for i, c in enumerate(cards):
        x = Inches(0.8) + i * Inches(4.0)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y_pos, col_w, col_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(0.25)

        p = ctf.paragraphs[0]
        p.text = c["num"]
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = c_blue

        p2 = ctf.add_paragraph()
        p2.text = c["title"]
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = c_navy
        p2.space_before = Pt(4)

        p3 = ctf.add_paragraph()
        p3.text = c["subtitle"]
        p3.font.size = Pt(11)
        p3.font.color.rgb = c_blue
        p3.space_before = Pt(2)

        for pt in c["points"]:
            pp = ctf.add_paragraph()
            pp.text = "• " + pt
            pp.font.size = Pt(11)
            pp.font.color.rgb = c_slate
            pp.space_before = Pt(10)

    # ==========================================
    # SLIDE 3: OUR SOLUTION & CORE FEATURES
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "The Solution: CheapPhone Lab Core Capabilities")

    feats = [
        {"icon": "⚡", "title": "Dual-Mode Testing", "desc": "Mode A: Deterministic zero-device simulation engine. Mode B: Real ADB emulator & Logcat integration."},
        {"icon": "📱", "title": "Indian Budget Profiles", "desc": "Pre-configured hardware profiles modeling 1GB-2GB RAM margins, Android Go editions, and tight flash storage."},
        {"icon": "🌐", "title": "Realistic Network Shaping", "desc": "Presets for 2G/EDGE (128kbps), high latency (RTT > 1500ms), elevator disconnects, and packet jitter."},
        {"icon": "⚠️", "title": "12 Failure Categories", "desc": "Detects Out-of-Memory crashes, ANR looper stalls, socket timeouts, lost activity state, and double submissions."},
        {"icon": "📊", "title": "Transparent Impact Model", "desc": "Quantifies affected users based on condition prevalence and target population with P0-P3 priority scoring."},
        {"icon": "🕸️", "title": "Neo4j Graph Explorer", "desc": "Graph pattern matching across Application -> Feature <- Failure <- TestRun -> Condition -> UserSegment."}
    ]

    for i, f in enumerate(feats):
        col = i % 3
        row = i // 3
        x = Inches(0.8) + col * Inches(4.0)
        y = Inches(1.6) + row * Inches(2.6)

        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.2)
        ctf.margin_right = Inches(0.2)
        ctf.margin_top = Inches(0.2)

        p = ctf.paragraphs[0]
        p.text = f"{f['icon']}  {f['title']}"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = c_navy

        p2 = ctf.add_paragraph()
        p2.text = f["desc"]
        p2.font.size = Pt(11)
        p2.font.color.rgb = c_slate
        p2.space_before = Pt(8)

    # ==========================================
    # SLIDE 4: END-TO-END PROJECT FLOW
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Project Workflow: From Configuration to Hotfix Recipe")

    steps = [
        {"step": "Step 1", "title": "Configure Test Lab", "desc": "Select budget profile, network preset (2G/3G), storage margin, memory pressure, and battery saver."},
        {"step": "Step 2", "title": "Execute Scenario", "desc": "Run deterministic simulation rules (instant) or launch controlled ADB Logcat on connected emulator."},
        {"step": "Step 3", "title": "Discover Failures", "desc": "App isolates OOM stacktraces, ANRs, socket timeouts, and state loss across affected app features."},
        {"step": "Step 4", "title": "Impact Triage", "desc": "Calculate estimated user reach: Population × Condition Prevalence × Failure Probability."},
        {"step": "Step 5", "title": "Neo4j Graph Analysis", "desc": "Traverse cross-run patterns to identify systemic hardware constraints causing correlated crashes."},
        {"step": "Step 6", "title": "Export Reproduction", "desc": "Download actionable step-by-step reproduction recipes in Markdown and JSON for engineering sprint hotfixes."}
    ]

    for i, st in enumerate(steps):
        col = i % 3
        row = i // 3
        x = Inches(0.8) + col * Inches(4.0)
        y = Inches(1.6) + row * Inches(2.6)

        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(0.2)

        p = ctf.paragraphs[0]
        p.text = st["step"].upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = c_blue

        p2 = ctf.add_paragraph()
        p2.text = st["title"]
        p2.font.size = Pt(15)
        p2.font.bold = True
        p2.font.color.rgb = c_navy
        p2.space_before = Pt(3)

        p3 = ctf.add_paragraph()
        p3.text = st["desc"]
        p3.font.size = Pt(11)
        p3.font.color.rgb = c_slate
        p3.space_before = Pt(8)

    # ==========================================
    # SLIDE 5: PROTOTYPE SHOWCASE - DASHBOARD
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Prototype Walkthrough: Central Dashboard")

    # Add screenshot on left
    img_path = "screenshots/dashboard.png"
    if os.path.exists(img_path):
        s5.shapes.add_picture(img_path, Inches(0.8), Inches(1.5), width=Inches(7.8))

    # Add description on right
    r_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(1.5), Inches(3.6), Inches(5.3))
    r_card.fill.solid()
    r_card.fill.fore_color.rgb = c_card_bg
    r_card.line.color.rgb = c_border
    rtf = r_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.25)
    rtf.margin_right = Inches(0.25)
    rtf.margin_top = Inches(0.3)

    p = rtf.paragraphs[0]
    p.text = "Executive Diagnostics"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_navy

    points = [
        "Live Run Metrics: Instant visibility into total test runs, pass/fail ratio, and critical P0/P1 blockers.",
        "Interactive Chart.js Visualizations: Failure distribution by category and severity triage breakdown.",
        "Priority Triage Queue: Direct highlights of top blocking failures and most vulnerable application modules.",
        "Recent Execution Activity: Real-time telemetry log tracking simulated and real ADB test outcomes."
    ]
    for pt in points:
        pp = rtf.add_paragraph()
        pp.text = "• " + pt
        pp.font.size = Pt(11)
        pp.font.color.rgb = c_slate
        pp.space_before = Pt(12)

    # ==========================================
    # SLIDE 6: PROTOTYPE SHOWCASE - TEST LAB
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Prototype Walkthrough: Test Lab Configuration")

    img_path = "screenshots/test_lab.png"
    if os.path.exists(img_path):
        s6.shapes.add_picture(img_path, Inches(0.8), Inches(1.5), width=Inches(7.8))

    r_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(1.5), Inches(3.6), Inches(5.3))
    r_card.fill.solid()
    r_card.fill.fore_color.rgb = c_card_bg
    r_card.line.color.rgb = c_border
    rtf = r_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.25)
    rtf.margin_right = Inches(0.25)
    rtf.margin_top = Inches(0.3)

    p = rtf.paragraphs[0]
    p.text = "Stress Simulation Engine"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_navy

    points = [
        "Mode A (Zero-Device Simulation): Immediate deterministic stress modeling for rapid hackathon evaluations.",
        "Mode B (Real ADB / Emulator): Auto-detects connected emulators and retrieves actual device properties via getprop.",
        "Configurable Stress Factors: Adjust network latency (2G/EDGE), available storage (<500MB), RAM pressure, and battery saver.",
        "Activity Interruption: Models background phone calls and Android Go OS task evictions."
    ]
    for pt in points:
        pp = rtf.add_paragraph()
        pp.text = "• " + pt
        pp.font.size = Pt(11)
        pp.font.color.rgb = c_slate
        pp.space_before = Pt(12)

    # ==========================================
    # SLIDE 7: PROTOTYPE SHOWCASE - IMPACT ANALYSIS
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Prototype Walkthrough: Transparent Impact Analysis")

    img_path = "screenshots/impact_analysis.png"
    if os.path.exists(img_path):
        s7.shapes.add_picture(img_path, Inches(0.8), Inches(1.5), width=Inches(7.8))

    r_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(1.5), Inches(3.6), Inches(5.3))
    r_card.fill.solid()
    r_card.fill.fore_color.rgb = c_card_bg
    r_card.line.color.rgb = c_border
    rtf = r_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.25)
    rtf.margin_right = Inches(0.25)
    rtf.margin_top = Inches(0.3)

    p = rtf.paragraphs[0]
    p.text = "Data-Driven Triage"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_navy

    points = [
        "Transparent Mathematical Model:\nAffected = Target Pop × Condition Prevalence × Failure Probability.",
        "Configurable Demographics: Customize MAU size, budget hardware percentage, and intermittent network coverage.",
        "Normalized Priority Scoring (0-100):\nSeverity (40%) + Reach (30%) + Repro (20%) + Confidence (10%).",
        "Actionable Classification: Categorizes failures into P0 Hotfix Blockers, P1 Sprint Priorities, and P2 Polish items."
    ]
    for pt in points:
        pp = rtf.add_paragraph()
        pp.text = "• " + pt
        pp.font.size = Pt(11)
        pp.font.color.rgb = c_slate
        pp.space_before = Pt(12)

    # ==========================================
    # SLIDE 8: PROTOTYPE SHOWCASE - NEO4J GRAPH
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Neo4j Graph Explorer: Cross-Run Pattern Matching")

    img_path = "screenshots/graph_explorer.png"
    if os.path.exists(img_path):
        s8.shapes.add_picture(img_path, Inches(0.8), Inches(1.5), width=Inches(7.8))

    r_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(1.5), Inches(3.6), Inches(5.3))
    r_card.fill.solid()
    r_card.fill.fore_color.rgb = c_card_bg
    r_card.line.color.rgb = c_border
    rtf = r_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.25)
    rtf.margin_right = Inches(0.25)
    rtf.margin_top = Inches(0.3)

    p = rtf.paragraphs[0]
    p.text = "Graph Traversal Value"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = c_navy

    points = [
        "Multi-Hop Schema: Connects Application -> AppFeature <- Failure <- TestRun -> Condition -> UserSegment.",
        "Systemic Root-Cause Discovery: Uncovers how disparate failures (e.g. OOM in Catalog & ANR on Startup) share identical stress nodes (2GB RAM).",
        "Interactive Vis-Network UI: Real-time visual graph canvas with drag, zoom, and neighborhood exploration.",
        "AuraDB Live / Fallback Store: Seamless Cypher execution with disconnected local storage fallback."
    ]
    for pt in points:
        pp = rtf.add_paragraph()
        pp.text = "• " + pt
        pp.font.size = Pt(11)
        pp.font.color.rgb = c_slate
        pp.space_before = Pt(12)

    # ==========================================
    # SLIDE 9: TECHNICAL ARCHITECTURE & STACK
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Technical Stack & Strict Architectural Adherence")

    tech_boxes = [
        {"layer": "Frontend Layer", "tech": "HTML5, CSS3, Vanilla JS, Bootstrap 5", "details": "Clean light developer dashboard, responsive layouts, zero heavy node frameworks, Chart.js for charts, Vis-network for graph visualization."},
        {"layer": "Backend & Core", "tech": "Python 3.11, Flask, python-dotenv", "details": "Modular architecture: core/test_runner.py, impact_analyzer.py, demo_scenarios.py, failure_detector.py with full REST API endpoints."},
        {"layer": "Graph Database", "tech": "Neo4j AuraDB & Official Python Driver", "details": "Real parameterized Cypher MERGE queries, schema constraints, subgraph extraction, and graceful disconnected demo store fallback."},
        {"layer": "Diagnostics & Testing", "tech": "ADB Subprocess, Logcat, pytest", "details": "Safe argument-list ADB runner (no shell injection), Logcat regex parsing for fatal crashes & ANRs, 5 automated pytest suites passing."}
    ]

    for i, t in enumerate(tech_boxes):
        y = Inches(1.6) + i * Inches(1.35)
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(1.15))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_border

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.15)

        p = ctf.paragraphs[0]
        p.text = f"{t['layer']}  —  {t['tech']}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = c_navy

        p2 = ctf.add_paragraph()
        p2.text = t["details"]
        p2.font.size = Pt(11)
        p2.font.color.rgb = c_slate
        p2.space_before = Pt(4)

    # ==========================================
    # SLIDE 10: CONCLUSION & IMPACT
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)

    # Decorative left accent bar
    bar = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.4), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = c_blue
    bar.line.fill.background()

    tbox = s10.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(10.5), Inches(4.5))
    tf10 = tbox.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.text = "CheapPhone Lab: Test Beyond the Flagship"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = c_navy

    p2 = tf10.add_paragraph()
    p2.text = "Empowering engineering teams to build software that thrives across India's next 500 million mobile users."
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = c_blue
    p2.space_before = Pt(10)

    recap_points = [
        "Discovers hidden failure modes before they reach play store reviews and cause user churn.",
        "Provides reproducible, step-by-step engineering hotfix recipes with Logcat stack traces.",
        "Demonstrates real Neo4j graph traversal value in debugging complex mobile system constraints.",
        "Works out of the box with zero configuration for hackathon evaluations and developer CI/CD."
    ]
    for pt in recap_points:
        pp = tf10.add_paragraph()
        pp.text = "✔  " + pt
        pp.font.size = Pt(14)
        pp.font.color.rgb = c_slate
        pp.space_before = Pt(14)

    output_path = "CheapPhone_Lab_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully as {output_path}")

if __name__ == "__main__":
    create_deck()
